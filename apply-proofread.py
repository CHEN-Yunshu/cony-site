#!/usr/bin/env python3
"""
把 Cony 校对完的 Excel 装回网站文案，未列入校对的语言保持不变。

用法: python3 apply-proofread.py 校对表.xlsx [--dry-run]

新版读取「文案」「交互提示」及表内「版本基准」，只导入实际改过的格子。
根据第 5 行表头识别语言列，兼容三语言及原四语言校对表。
网站和 Excel 同时改过同一格时会报告冲突并跳过，避免覆盖新内容。
旧表仍与 content/proofread-baseline.json 比较，沿用原有导入流程。
空白格视为未修改；要删除文字，请把整格改成［删除］。
写入前备份修改到的 i18n.js / app.js；--dry-run 不写入或备份。
"""
import argparse
import html
import json
from pathlib import Path
import re
import shutil
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
LANGS = ('en', 'zh', 'ko', 'th')
SOURCE_FILES = {'i18n': 'i18n.js', 'ui': 'app.js'}
JS_STRING = r"'((?:[^'\\\r\n]|\\(?:\r\n|[\s\S]))*)'"
ENTRY_PATTERN = re.compile(
    r'^[ \t]+(?P<key>[A-Za-z_$][\w$]*):[ \t]*\{\s*'
    + r'\s*'.join(rf'{lg}:\s*{JS_STRING},' for lg in LANGS)
    + r'\s*\},?', re.M)
UI_ENTRY_PATTERN = re.compile(
    r'^[ \t]+(?P<key>[A-Za-z_$][\w$]*):[ \t]*\[[ \t]*'
    + r'[ \t]*,[ \t]*'.join([JS_STRING] * 4)
    + r'[ \t]*\],[ \t]*,?[ \t]*$', re.M)


def to_html(value):
    """Restore the workbook's editorial markers; keep legacy conversion rules."""
    if value is None:
        return None
    value = str(value)
    value = (value.replace('&', '&amp;').replace('"', '&quot;')
             .replace('［空行］', '<br><br>').replace('［换行］', '<br>')
             .replace('［粗体］', '<b>').replace('［/粗体］', '</b>'))
    return value.replace('\r', '').strip('\n')


def display_html(value):
    """Match the readable text exported to the new workbook."""
    value = re.sub(r'<br\s*/?>', '［换行］', value, flags=re.I)
    value = value.replace('<b>', '［粗体］').replace('</b>', '［/粗体］')
    return html.unescape(value)


def esc_js(value):
    escapes = {'\\': '\\\\', "'": "\\'", '\n': '\\n', '\r': '\\r',
               '\t': '\\t', '\b': '\\b', '\f': '\\f', '\v': '\\v'}
    return ''.join(escapes.get(ch, '\\u%04x' % ord(ch)
                   if ord(ch) < 32 or ord(ch) in (0x2028, 0x2029)
                   or 0xD800 <= ord(ch) <= 0xDFFF else ch) for ch in value)


def unesc_js(value):
    """Decode a JS single-quoted string literal without executing JavaScript."""
    output = []
    index = 0
    simple = {'n': '\n', 'r': '\r', 't': '\t', 'b': '\b', 'f': '\f',
              'v': '\v', '0': '\0'}
    while index < len(value):
        ch = value[index]
        index += 1
        if ch != '\\':
            output.append(ch)
            continue
        if index == len(value):
            raise ValueError('unterminated JavaScript escape')
        ch = value[index]
        index += 1
        if ch in '\r\n':
            if ch == '\r' and value[index:index + 1] == '\n':
                index += 1
            continue
        if ch in ('u', 'x'):
            if ch == 'u' and value[index:index + 1] == '{':
                end = value.find('}', index + 1)
                digits = value[index + 1:end] if end >= 0 else ''
                index = end + 1
            else:
                length = 4 if ch == 'u' else 2
                digits = value[index:index + length]
                if len(digits) != length:
                    raise ValueError('incomplete JavaScript Unicode escape')
                index += length
            if not re.fullmatch(r'[0-9a-fA-F]+', digits):
                raise ValueError('invalid JavaScript Unicode escape')
            output.append(chr(int(digits, 16)))
        elif ch.isdigit() and (ch != '0' or value[index:index + 1].isdigit()):
            raise ValueError('legacy octal JavaScript escapes are not supported')
        else:
            output.append(simple.get(ch, ch))
    # JS may spell an emoji as a pair of UTF-16 escapes.
    return ''.join(output).encode('utf-16-le', 'surrogatepass').decode('utf-16-le', 'surrogatepass')


def parse_source(text, source):
    """Return decoded strings and exact replacement spans, never eval JS."""
    if source == 'ui':
        block = re.search(r'\bconst\s+ui\s*=\s*\{(?P<body>.*?)^[ \t]*\};', text, re.M | re.S)
        if not block:
            raise ValueError('app.js: cannot find the const ui object')
        fragment, offset = block.group('body'), block.start('body')
        pattern = UI_ENTRY_PATTERN
    else:
        fragment, offset = text, 0
        pattern = ENTRY_PATTERN
    result = {}
    for match in pattern.finditer(fragment):
        key = ('ui.' if source == 'ui' else '') + match.group('key')
        if key in result:
            raise ValueError(f'{SOURCE_FILES[source]}: duplicate key {key}')
        result[key] = {
            'values': {lg: unesc_js(match.group(i + 2)) for i, lg in enumerate(LANGS)},
            'spans': {lg: (offset + match.start(i + 2), offset + match.end(i + 2))
                      for i, lg in enumerate(LANGS)},
        }
    if not result:
        raise ValueError(f'{SOURCE_FILES[source]}: no supported four-language entries found')
    return result


def cell_value(row, column):
    if column >= len(row):
        return None
    cell = row[column]
    if cell.data_type == 'f':
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: formulas are not accepted as copy')
    if cell.data_type == 'e':
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: spreadsheet error {cell.value}')
    return cell.value


def language_columns(sheet, header_row):
    """Read language headers, so a shifted status column can never become copy."""
    headers = next(sheet.iter_rows(min_row=header_row, max_row=header_row))
    columns = {}
    aliases = {'en': ('english', 'en'), 'zh': ('中文', 'zh'),
               'ko': ('한국어', 'ko'), 'th': ('ไทย', 'th')}
    for index, cell in enumerate(headers):
        label = str(cell_value(headers, index) or '').strip().lower()
        for language, names in aliases.items():
            if any(label == name or label.startswith(name + '（')
                   or label.startswith(name + ' (') for name in names):
                if language in columns:
                    raise ValueError(f'{sheet.title}: duplicate language column {language}')
                columns[language] = index
    if not all(language in columns for language in ('en', 'zh', 'ko')):
        raise ValueError(f'{sheet.title}: English, 中文 and 한국어 headers are required')
    return columns


def load_edits(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        modern = '版本基准' in workbook.sheetnames
        baseline, sources = {}, {}
        if modern:
            for row in workbook['版本基准'].iter_rows(min_row=2):
                key = cell_value(row, 0)
                if key is None or not str(key).strip():
                    continue
                key = str(key).strip()
                source = cell_value(row, 1)
                if source not in SOURCE_FILES or (source == 'ui') != key.startswith('ui.'):
                    raise ValueError(f'版本基准: invalid source/key pair {source!r}, {key!r}')
                if key in baseline:
                    raise ValueError(f'版本基准: duplicate key {key}')
                values = [cell_value(row, i + 2) for i in range(4)]
                if any(v is not None and not isinstance(v, str) for v in values):
                    raise ValueError(f'版本基准: non-text original value for {key}')
                baseline[key] = dict(zip(LANGS, (v if v is not None else '' for v in values)))
                sources[key] = source
            if not baseline:
                raise ValueError('版本基准 is empty; refusing to import without originals')
        else:
            # A missing legacy baseline must never cause a blind overwrite.
            baseline = json.loads((ROOT / 'content/proofread-baseline.json').read_text(encoding='utf-8'))
            sources = dict.fromkeys(baseline, 'i18n')

        edits = {}
        sheets = ('文案', '交互提示') if modern else ('文案',)
        for name in sheets:
            if name not in workbook.sheetnames:
                raise ValueError(f'missing worksheet: {name}')
            columns = (language_columns(workbook[name], 5) if modern
                       else {language: index + 2 for index, language in enumerate(LANGS)})
            for row in workbook[name].iter_rows(min_row=6 if modern else 2):
                key = cell_value(row, 0)
                if key is None or not str(key).strip() or str(key).startswith('▌'):
                    continue
                key = str(key).strip()
                if key in edits:
                    raise ValueError(f'duplicate worksheet key: {key}')
                expected_source = 'ui' if name == '交互提示' else 'i18n'
                if key in sources and sources[key] != expected_source:
                    raise ValueError(f'{name}: {key} belongs in a different worksheet')
                values = {}
                for language, index in columns.items():
                    value = cell_value(row, index)
                    if value is None or not str(value).strip():
                        values[language] = None
                    elif str(value).strip() == '［删除］':
                        values[language] = ''
                    else:
                        values[language] = (to_html(value) if expected_source == 'i18n'
                                            else str(value).replace('\r\n', '\n').replace('\r', '\n'))
                edits[key] = values
        return edits, baseline, sources, modern
    finally:
        workbook.close()


def plan_changes(edits, baseline, sources, parsed, modern):
    replacements = {source: [] for source in parsed}
    changed, conflicts, unknown, missing = [], [], [], []
    for key, new in edits.items():
        if key not in baseline:
            unknown.append(key)
            continue
        source = sources[key]
        if key not in parsed.get(source, {}):
            missing.append(key)
            continue
        entry = parsed[source][key]
        original = baseline[key]
        for language in LANGS:
            revised = new.get(language)
            if revised is None or language not in original:
                continue
            old = original[language]
            # Entity spelling may normalize during export (e.g. &#39; -> ').
            unchanged = {old}
            if modern and source == 'i18n':
                unchanged.add(to_html(display_html(old)))
            current = entry['values'][language]
            if revised in unchanged or revised == current:
                continue
            if modern and current != old:
                conflicts.append((key, language, old, current, revised))
                continue
            start, end = entry['spans'][language]
            replacements[source].append((start, end, esc_js(revised)))
            changed.append((source, key, language, current, revised))
    return replacements, changed, conflicts, unknown, missing


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('xlsx', nargs='?', default=ROOT / '文案校对表_Cony网站.xlsx')
    parser.add_argument('--dry-run', action='store_true', help='只报告修改，不写入或备份')
    args = parser.parse_args()

    try:
        edits, baseline, sources, modern = load_edits(args.xlsx)
        active_sources = ('i18n', 'ui') if modern else ('i18n',)
        texts = {source: (ROOT / SOURCE_FILES[source]).read_text(encoding='utf-8') for source in active_sources}
        parsed = {source: parse_source(text, source) for source, text in texts.items()}
        replacements, changed, conflicts, unknown, missing = plan_changes(
            edits, baseline, sources, parsed, modern)
    except (OSError, KeyError, ValueError) as error:
        print(f'error: {error}', file=sys.stderr)
        return 2

    print(f'read {len(edits)} keys from {args.xlsx}')
    print('baseline: ' + ('版本基准 (embedded; conflict protection enabled)' if modern
                         else 'content/proofread-baseline.json (legacy workflow)'))
    for source, entries in parsed.items():
        print(f'matched {len(entries)} entries in {SOURCE_FILES[source]}')
    if unknown:
        print(f'!! unknown sheet keys without a baseline (ignored): {sorted(unknown)}')
    if missing:
        print(f'!! keys absent from the supported current source format (not inserted): {sorted(missing)}')
    for key, language, old, current, revised in conflicts:
        print(f'!! CONFLICT {key} [{language}] — skipped\n'
              f'   exported: {old!r}\n   website:  {current!r}\n   Excel:    {revised!r}')

    if args.dry_run:
        print(f'\ndry-run: {len(changed)} change(s), {len(conflicts)} conflict(s); source files and backups untouched')
    elif changed:
        for source, spans in replacements.items():
            if not spans:
                continue
            new_text = texts[source]
            for start, end, value in sorted(spans, reverse=True):
                new_text = new_text[:start] + value + new_text[end:]
            source_path = ROOT / SOURCE_FILES[source]
            shutil.copy2(source_path, source_path.with_suffix('.js.bak'))
            source_path.write_text(new_text, encoding='utf-8')
        print(f'\napplied {len(changed)} change(s); skipped {len(conflicts)} conflict(s)')
    else:
        print(f'no changes — source files untouched; {len(conflicts)} conflict(s)')
    for source, key, language, old, new in changed:
        print(f'  {SOURCE_FILES[source]} / {key} [{language}]\n'
              f'    {old[:100]!r}\n    -> {new[:100]!r}')
    if changed and not args.dry_run:
        print('\nModified sources were backed up as .js.bak — now run ./build-preview.sh')
    return 1 if conflicts or unknown or missing else 0


if __name__ == '__main__':
    sys.exit(main())
